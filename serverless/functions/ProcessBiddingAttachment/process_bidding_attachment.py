import io, logging, mimetypes, pathlib, tempfile, typing as t

import docx, magic, openpyxl, pypdf, py7zr, textract, textract.exceptions, zipfile

import nltk, nltk.corpus, nltk.tokenize, nltk.stem

nltk.data.path.append('.nltk')

nltk.download("punkt_tab", download_dir=".nltk")
nltk.download("stopwords", download_dir=".nltk")

from petronect.persistence.Persistence import Persistence
from petronect.persistence.PetronectAttachmentPersistence import PetronectAttachmentPersistence
from petronect.persistence.PetronectBiddingPersistence import PetronectBiddingPersistence

from serverless.FunctionInvoker import FunctionInvoker
from serverless.KeyValueStorage import KeyValueStorage


logger = logging.getLogger()

logger.setLevel(logging.INFO)

def lambda_handler(event, context):
    assert isinstance(event, dict)
    assert "storage_id" in event

    bidding_hash, file_name = event['storage_id'].split(":", 1)

    PetronectBiddingPersistence.lock(bidding_hash)

    logger.info(f"Processing: {bidding_hash} {file_name}")

    storage = KeyValueStorage()

    data = storage.load(event['storage_id'])

    tokens = extract_tokens(file_name=file_name, buf=data)

    if isinstance(tokens, set):
        PetronectAttachmentPersistence.create(
            bidding_hash=bidding_hash,
            name=file_name,
            tokens=tokens
        )
    else:
        for k, v in tokens["extraction"].items():
            PetronectAttachmentPersistence.create(
                bidding_hash=bidding_hash,
                name=k,
                tokens=v
            )

    PetronectBiddingPersistence.unlock(bidding_hash)

    invoker = FunctionInvoker()
    invoker.trigger(
        "match_bidding",
        dict(bidding_hash=bidding_hash)
    )

    Persistence.commit()

def extract_tokens(file_name: str, buf: bytes) -> t.Union[set[str], dict[str, dict]]:
    ext = get_extension(file_name=file_name, buf=buf)

    if ext is None:
        raise Exception(f"Unsupported file: {file_name}.")

    if ext in [".zip", ".7z"]:
        return extract_tokens_from_compressed_archive(buf, ext)

    text = extract_text(buf, ext=ext)

    if text is None:
        raise Exception(f"Text extraction failed for file: {file_name}.")

    return tokenize_text(text)

def get_extension(file_name: str, buf: bytes) -> t.Optional[str]:
    suffix = pathlib.Path(file_name).suffix.lower()

    if suffix and mimetypes.types_map.get(suffix):
        return suffix

    mime_type = get_mime_type(buf)
    guessed = mimetypes.guess_extension(mime_type)

    return guessed

def get_mime_type(buf: bytes) -> str:
    return magic.Magic(mime=True).from_buffer(buf)

def extract_tokens_from_compressed_archive(buf: bytes, ext: str) -> dict:
    result = dict(ERRORS=dict(), extraction=dict())

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir_path = pathlib.Path(tmpdir)

        archive_path = tmpdir_path / f"archive{ext}"
        archive_path.write_bytes(buf)

        extract_dir = tmpdir_path / "extracted"
        extract_dir.mkdir()

        if ext == ".zip":
            with zipfile.ZipFile(archive_path, "r") as archive:
                archive.extractall(extract_dir)
        elif ext == ".7z":
            with py7zr.SevenZipFile(archive_path, "r") as archive:
                archive.extractall(path=extract_dir)

        for inner_path in extract_dir.rglob("*"):
            if inner_path.is_file():
                relative_path = inner_path.relative_to(extract_dir).as_posix()

                inner_bytes = inner_path.read_bytes()

                inner_ext = get_extension(inner_path.name, inner_bytes)

                if inner_ext is None:
                    result["ERRORS"][relative_path] = f"Could not infer file extension: {inner_path.name}"
                    continue

                text = extract_text(inner_bytes, ext=inner_ext)

                if text is None:
                    result["ERRORS"][relative_path] = f"Could not extract text from file : {relative_path}"
                    continue

                result["extraction"][relative_path] = tokenize_text(text)

    return result

def tokenize_text(text: str) -> set[str]:
    stop_words = set(nltk.corpus.stopwords.words("portuguese"))
    stemmer = nltk.stem.SnowballStemmer("portuguese")

    tokens = nltk.tokenize.word_tokenize(text.lower(), language="portuguese")
    filtered = [ stemmer.stem(t) for t in tokens if t.isalpha() and t not in stop_words ]

    return set(filtered)

def extract_text(buf: bytes, ext: str) -> t.Optional[str]:
    if ext == ".txt":
        return extract_text_txt(buf)
    elif ext == ".pdf":
        return extract_text_pdf(buf)
    elif ext == ".docx":
        return extract_text_docx(buf)
    elif ext == ".doc":
        try:
            return extract_text_doc(buf)
        except textract.exceptions.ShellError:
            return None
    elif ext == ".xlsx":
        return extract_text_xlsx(buf)
    else:
        raise Exception(f"Unsupported file: {ext}.")

def extract_text_txt(buf: bytes) -> str:
    return buf.decode(errors='ignore')

def extract_text_pdf(buf: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(buf))
    return "\n".join(page.extract_text() or '' for page in reader.pages)

def extract_text_docx(buf: bytes) -> str:
    doc = docx.Document(io.BytesIO(buf))
    return "\n".join(p.text for p in doc.paragraphs)

def extract_text_doc(buf: bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=True, suffix=".doc") as temp_file:
        temp_file.write(buf)
        temp_file.flush()

        return textract.process(temp_file.name, encoding='utf-8').decode("utf-8", errors="ignore")

def extract_text_xlsx(buf: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(buf), read_only=True, data_only=True)

    output = list()

    for sheet in wb.worksheets:
        output.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):

            row_text = "\t".join(str(cell) for cell in row if cell is not None)

            if row_text.strip():
                output.append(row_text)

    return "\n".join(output)
